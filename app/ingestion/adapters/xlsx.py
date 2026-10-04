"""XLSX adapter: sheet selection, header mapping, and multiple sheets (records / rules / terms).

Mapping:
    {"records_sheet": "Codes",          # default: first sheet
     "header_row": 1,
     "rules_sheet": "Rules",            # optional long-format rules sheet
     "index_terms_sheet": "Index",      # optional long-format term sheet
     "columns": {...}, "rule_columns": {...}, "list_separator": "|"}
"""

from collections.abc import Iterator
from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.constants import SourceType
from app.ingestion.adapters.base import AdapterError, SourceAdapter, SourceManifest
from app.ingestion.adapters.fields import (
    FieldMapping,
    RecordAssembler,
    canonical_row,
    canonical_rule_row,
)
from app.ingestion.models import NormalizedICDRecord, SourceProvenance


class XlsxMapping(FieldMapping):
    records_sheet: str | None = None
    header_row: int = 1
    rules_sheet: str | None = None
    index_terms_sheet: str | None = None


class XlsxAdapter(SourceAdapter):
    source_type: ClassVar[SourceType] = SourceType.XLSX
    extensions: ClassVar[tuple[str, ...]] = (".xlsx", ".xlsm")

    def __init__(self, path: Path, manifest: SourceManifest | None = None) -> None:
        super().__init__(path, manifest)
        self.mapping = XlsxMapping.model_validate(self.manifest.mapping)

    @classmethod
    def detect(cls, path: Path) -> float:
        if path.suffix.lower() not in cls.extensions:
            return 0.0
        with path.open("rb") as handle:
            return 0.9 if handle.read(4) == b"PK\x03\x04" else 0.1

    def _workbook(self):  # noqa: ANN202 - openpyxl Workbook
        try:
            # read_only streams rows; data_only returns cached values instead of formulas.
            return load_workbook(self.path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, KeyError, ValueError) as exc:
            raise AdapterError(f"Cannot open workbook {self.path.name}: {exc}") from exc

    def _sheet_rows(self, workbook, sheet_name: str) -> Iterator[tuple[int, dict[str, Any]]]:  # noqa: ANN001
        if sheet_name not in workbook.sheetnames:
            raise AdapterError(f"Sheet {sheet_name!r} not found in {self.path.name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header: list[str] | None = None
        for row_index, values in enumerate(rows, start=1):
            if row_index < self.mapping.header_row:
                continue
            if row_index == self.mapping.header_row:
                header = [str(v).strip() if v is not None else "" for v in values]
                continue
            if header is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            yield row_index, {h: v for h, v in zip(header, values, strict=False) if h}

    def _inspection_details(self) -> dict[str, Any]:
        workbook = self._workbook()
        try:
            return {
                "sheets": {name: workbook[name].max_row for name in workbook.sheetnames},
            }
        finally:
            workbook.close()

    def iterate_records(self) -> Iterator[NormalizedICDRecord]:
        workbook = self._workbook()
        try:
            records_sheet = self.mapping.records_sheet or workbook.sheetnames[0]
            assembler = RecordAssembler(self.mapping)
            for row_index, row in self._sheet_rows(workbook, records_sheet):
                provenance = SourceProvenance(
                    source_filename=self.path.name,
                    kind=self.source_type,
                    sheet=records_sheet,
                    row=row_index,
                )
                assembler.add_row(canonical_row(row, self.mapping), provenance)
            for sheet in (self.mapping.rules_sheet, self.mapping.index_terms_sheet):
                if not sheet:
                    continue
                for row_index, row in self._sheet_rows(workbook, sheet):
                    provenance = SourceProvenance(
                        source_filename=self.path.name,
                        kind=self.source_type,
                        sheet=sheet,
                        row=row_index,
                    )
                    assembler.add_rule_row(canonical_rule_row(row, self.mapping), provenance)
        finally:
            workbook.close()
        self._issues.extend(assembler.issues)
        yield from assembler.records()
